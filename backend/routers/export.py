import json
import os
import io
from fastapi import APIRouter, HTTPException, Response
from sqlmodel import Session, select
from datetime import datetime

from models import AnalysisRun, AttackChain, Finding, engine

router = APIRouter(prefix="/api/runs", tags=["export"])

@router.get("/{run_id}/export")
async def export_run(run_id: int, format: str = "json"):
    with Session(engine) as session:
        run = session.get(AnalysisRun, run_id)
        if not run:
            raise HTTPException(status_code=404, detail="Run not found.")

        findings = session.exec(select(Finding).where(Finding.run_id == run_id)).all()
        chains = session.exec(select(AttackChain).where(AttackChain.run_id == run_id).order_by(AttackChain.chain_index)).all()

        report_data = {
            "run_id": run.id,
            "repo_url": run.repo_url,
            "status": run.status,
            "created_at": run.created_at.isoformat(),
            "findings_count": len(findings),
            "chains": [
                {
                    "title": c.title,
                    "impact": c.impact,
                    "confidence": c.confidence,
                    "steps": json.loads(c.steps_json),
                    "explanation": c.explanation
                }
                for c in chains
            ],
            "findings": [
                {
                    "source": f.source,
                    "severity": f.severity,
                    "file_path": f.file_path,
                    "line_start": f.line_start,
                    "message": f.message
                }
                for f in findings
            ]
        }

        if format == "json":
            return Response(
                content=json.dumps(report_data, indent=2),
                media_type="application/json",
                headers={"Content-Disposition": f"attachment; filename=acm_report_{run_id}.json"}
            )
        elif format == "docx":
            try:
                from docx import Document
            except ImportError:
                raise HTTPException(status_code=500, detail="python-docx not installed.")
            
            doc = Document()
            doc.add_heading(f"Attack Chain Mapper Report - Run {run_id}", 0)
            doc.add_paragraph(f"Repository: {run.repo_url}")
            doc.add_paragraph(f"Date: {run.created_at.strftime('%Y-%m-%d %H:%M')}")
            
            doc.add_heading("Attack Chains", level=1)
            for chain in report_data["chains"]:
                p = doc.add_paragraph()
                p.add_run(chain['title']).bold = True
                p.add_run(f"\nImpact: {chain['impact']} | Confidence: {chain['confidence']}")
                doc.add_paragraph(chain['explanation'])
                
            doc.add_heading("Findings", level=1)
            for f in report_data["findings"]:
                doc.add_paragraph(f"[{f['severity']}] {f['file_path']}:{f['line_start']} - {f['source']}")
                doc.add_paragraph(f['message'])
                
            f = io.BytesIO()
            doc.save(f)
            f.seek(0)
            return Response(
                content=f.read(),
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f"attachment; filename=acm_report_{run_id}.docx"}
            )
        elif format == "pdf":
            try:
                from fpdf import FPDF
            except ImportError:
                raise HTTPException(status_code=500, detail="fpdf2 not installed.")
                
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("helvetica", size=16, style="B")
            pdf.cell(200, 10, f"Attack Chain Mapper Report - Run {run_id}", align='C')
            pdf.ln(10)
            pdf.set_font("helvetica", size=12)
            pdf.cell(200, 10, f"Repository: {run.repo_url}", align='C')
            pdf.ln(10)
            pdf.cell(200, 10, f"Date: {run.created_at.strftime('%Y-%m-%d %H:%M')}", align='C')
            pdf.ln(10)
            
            pdf.set_font("helvetica", size=14, style="B")
            pdf.cell(200, 10, "Attack Chains")
            pdf.ln(10)
            pdf.set_font("helvetica", size=11)
            for chain in report_data["chains"]:
                pdf.set_font("helvetica", style="B")
                pdf.multi_cell(0, 10, f"{chain['title']}")
                pdf.set_font("helvetica")
                pdf.multi_cell(0, 8, f"Impact: {chain['impact']} | Confidence: {chain['confidence']}")
                text = chain['explanation'].encode('latin-1', 'replace').decode('latin-1')
                pdf.multi_cell(0, 8, text)
                pdf.ln(5)
                
            pdf.add_page()
            pdf.set_font("helvetica", size=14, style="B")
            pdf.cell(200, 10, "Findings")
            pdf.ln(10)
            pdf.set_font("helvetica", size=11)
            for f in report_data["findings"]:
                pdf.multi_cell(0, 8, f"[{f['severity']}] {f['file_path']}:{f['line_start']} - {f['source']}")
                text = f['message'].encode('latin-1', 'replace').decode('latin-1')
                pdf.multi_cell(0, 8, text)
                pdf.ln(2)
                
            pdf_bytes = pdf.output()
            return Response(
                content=bytes(pdf_bytes),
                media_type="application/pdf",
                headers={"Content-Disposition": f"attachment; filename=acm_report_{run_id}.pdf"}
            )
        else:
            raise HTTPException(status_code=400, detail="Invalid format.")
